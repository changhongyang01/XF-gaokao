# -*- coding: utf-8 -*-
import sys, os, json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

# === 读取数据 ===
input_file = sys.argv[1] if len(sys.argv) > 1 else None
if not input_file or not os.path.exists(input_file):
    print("Usage: python generate_excel.py data.json")
    print("Error: data.json not found. Using sample data.")
    data = {
        "user": {"name": "示例", "province": "河南", "type": "美术类", "cultural": 350, "professional": 203, "comprehensive": 428.75},
        "targets": [
            {"school": "郑州科技学院", "major": "视觉传达", "tier": "冲", "score_2023": 423, "score_2024": 428, "score_2025": 432, "tuition": 16000, "pros": "省内认可度高", "cons": "分数线略高"},
            {"school": "郑州经贸学院", "major": "产品设计", "tier": "稳", "score_2023": 418, "score_2024": 423, "score_2025": 426, "tuition": 15000, "pros": "性价比高", "cons": "专业偏工科"},
            {"school": "河南科传学院", "major": "环境设计", "tier": "稳", "score_2023": 415, "score_2024": 420, "score_2025": 424, "tuition": 16000, "pros": "河大品牌", "cons": "新建校区"},
            {"school": "郑州商学院", "major": "数媒", "tier": "保", "score_2023": 412, "score_2024": 417, "score_2025": 420, "tuition": 15000, "pros": "稳录", "cons": "品牌较弱"},
        ]
    }
else:
    with open(input_file, 'r', encoding='utf-8-sig') as f:
        data = json.load(f)

user = data.get("user", {})
targets = data.get("targets", [])
comprehensive = user.get("comprehensive", 0)

# === 样式定义 ===
DARK_BLUE = "1F4E79"
LIGHT_RED = "FFE0E0"
LIGHT_YELLOW = "FFF8E0"
LIGHT_GREEN = "E0FFE0"
WHITE = "FFFFFF"
BORDER_COLOR = "B0B0B0"

header_fill = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type="solid")
header_font = Font(name="微软雅黑", size=11, bold=True, color=WHITE)
title_font = Font(name="微软雅黑", size=16, bold=True, color=DARK_BLUE)
subtitle_font = Font(name="微软雅黑", size=10, color="666666")
normal_font = Font(name="微软雅黑", size=10)
bold_font = Font(name="微软雅黑", size=10, bold=True)
red_font = Font(name="微软雅黑", size=10, bold=True, color="CC0000")
green_font = Font(name="微软雅黑", size=10, color="007A33")
red_fill = PatternFill(start_color=LIGHT_RED, end_color=LIGHT_RED, fill_type="solid")
yellow_fill = PatternFill(start_color=LIGHT_YELLOW, end_color=LIGHT_YELLOW, fill_type="solid")
green_fill = PatternFill(start_color=LIGHT_GREEN, end_color=LIGHT_GREEN, fill_type="solid")
thin_border = Border(
    left=Side(style='thin', color=BORDER_COLOR),
    right=Side(style='thin', color=BORDER_COLOR),
    top=Side(style='thin', color=BORDER_COLOR),
    bottom=Side(style='thin', color=BORDER_COLOR)
)
center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

def apply_cell(ws, row, col, value, font=normal_font, fill=None, alignment=center_align):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font
    cell.alignment = alignment
    cell.border = thin_border
    if fill:
        cell.fill = fill
    return cell

def auto_width(ws, min_width=10, max_width=30):
    for col_cells in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col_cells), default=0)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len * 2 + 2, min_width), max_width)

# === 创建工作簿 ===
wb = openpyxl.Workbook()

# ========== Sheet 1: 冲稳保总览 ==========
ws1 = wb.active
ws1.title = "冲稳保总览"

# 标题
ws1.merge_cells("A1:H1")
apply_cell(ws1, 1, 1, f"{user.get('province','')}{user.get('type','')}志愿填报方案 | 综合分: {comprehensive}", title_font, alignment=Alignment(horizontal="center", vertical="center"))
ws1.row_dimensions[1].height = 35

headers1 = ["层级", "院校", "推荐专业", "2023", "2024", "2025预估", "学费/年", "你的差距"]
for col, h in enumerate(headers1, 1):
    apply_cell(ws1, 2, col, h, header_font, header_fill)

for i, t in enumerate(targets):
    row = i + 3
    gap = comprehensive - t.get("score_2025", 0)
    tier = t.get("tier", "")
    
    row_fill = red_fill if tier == "冲" else (yellow_fill if tier == "稳" else green_fill)
    gap_font = red_font if gap < 0 else green_font
    
    apply_cell(ws1, row, 1, tier, bold_font, row_fill)
    apply_cell(ws1, row, 2, t["school"], normal_font, row_fill)
    apply_cell(ws1, row, 3, t["major"], normal_font, row_fill)
    apply_cell(ws1, row, 4, t.get("score_2023", "-"), normal_font, row_fill)
    apply_cell(ws1, row, 5, t.get("score_2024", "-"), normal_font, row_fill)
    apply_cell(ws1, row, 6, t.get("score_2025", "-"), normal_font, row_fill)
    apply_cell(ws1, row, 7, f"¥{t.get('tuition',0):,}", normal_font, row_fill)
    apply_cell(ws1, row, 8, f"{gap:+.1f}", gap_font, row_fill)

ws1.freeze_panes = "A3"
auto_width(ws1)
ws1.column_dimensions['B'].width = 18
ws1.column_dimensions['C'].width = 14

# ========== Sheet 2: 近三年趋势 ==========
ws2 = wb.create_sheet("近三年趋势")

ws2.merge_cells("A1:F1")
apply_cell(ws2, 1, 1, "近三年录取分数走势", title_font, alignment=Alignment(horizontal="center", vertical="center"))
ws2.row_dimensions[1].height = 35

headers2 = ["院校", "2023", "2024", "2025预估", "趋势", "安全评估"]
for col, h in enumerate(headers2, 1):
    apply_cell(ws2, 2, col, h, header_font, header_fill)

for i, t in enumerate(targets):
    row = i + 3
    s23, s24, s25 = t.get("score_2023", 0), t.get("score_2024", 0), t.get("score_2025", 0)
    trend = "⬆️上涨" if s25 > s24 else ("⬇️下降" if s25 < s24 else "➡️持平")
    gap = comprehensive - s25
    safety = "🟢安全" if gap > 5 else ("🟡接近" if gap > -3 else "🔴危险")
    
    tier = t.get("tier", "")
    row_fill = red_fill if tier == "冲" else (yellow_fill if tier == "稳" else green_fill)
    
    apply_cell(ws2, row, 1, t["school"], normal_font, row_fill)
    apply_cell(ws2, row, 2, s23 if s23 else "-", normal_font, row_fill)
    apply_cell(ws2, row, 3, s24 if s24 else "-", normal_font, row_fill)
    apply_cell(ws2, row, 4, s25 if s25 else "-", normal_font, row_fill)
    apply_cell(ws2, row, 5, trend, normal_font, row_fill)
    apply_cell(ws2, row, 6, safety, bold_font, row_fill)

ws2.freeze_panes = "A3"
auto_width(ws2)
ws2.column_dimensions['A'].width = 18

# ========== Sheet 3: 院校详情 ==========
ws3 = wb.create_sheet("院校详情")

ws3.merge_cells("A1:F1")
apply_cell(ws3, 1, 1, "各院校详细评估", title_font, alignment=Alignment(horizontal="center", vertical="center"))
ws3.row_dimensions[1].height = 35

headers3 = ["院校", "专业", "层级", "学费(元/年)", "核心优势", "注意事项"]
for col, h in enumerate(headers3, 1):
    apply_cell(ws3, 2, col, h, header_font, header_fill)

for i, t in enumerate(targets):
    row = i + 3
    tier = t.get("tier", "")
    row_fill = red_fill if tier == "冲" else (yellow_fill if tier == "稳" else green_fill)
    
    apply_cell(ws3, row, 1, t["school"], bold_font, row_fill)
    apply_cell(ws3, row, 2, t["major"], normal_font, row_fill)
    apply_cell(ws3, row, 3, tier, bold_font, row_fill)
    apply_cell(ws3, row, 4, f"¥{t.get('tuition',0):,}", normal_font, row_fill)
    apply_cell(ws3, row, 5, t.get("pros", ""), normal_font, row_fill, left_align)
    apply_cell(ws3, row, 6, t.get("cons", ""), normal_font, row_fill, left_align)

ws3.freeze_panes = "A3"
auto_width(ws3)
ws3.column_dimensions['A'].width = 18
ws3.column_dimensions['E'].width = 30
ws3.column_dimensions['F'].width = 30

# ========== Sheet 4: 费用概算 ==========
ws4 = wb.create_sheet("费用概算")

ws4.merge_cells("A1:F1")
apply_cell(ws4, 1, 1, "四年总费用概算", title_font, alignment=Alignment(horizontal="center", vertical="center"))
ws4.row_dimensions[1].height = 35

headers4 = ["院校", "学费/年", "住宿费/年", "生活费/年", "四年总计", "ROI评估"]
for col, h in enumerate(headers4, 1):
    apply_cell(ws4, 2, col, h, header_font, header_fill)

for i, t in enumerate(targets):
    row = i + 3
    tuition = t.get("tuition", 0)
    dorm = 1500
    living = 15000
    total = (tuition + dorm + living) * 4
    roi = "2-3年回本" if total < 120000 else ("3-4年回本" if total < 150000 else "4年以上回本")
    
    apply_cell(ws4, row, 1, t["school"], normal_font)
    apply_cell(ws4, row, 2, f"¥{tuition:,}", normal_font)
    apply_cell(ws4, row, 3, f"¥{dorm:,}", normal_font)
    apply_cell(ws4, row, 4, f"¥{living:,}", normal_font)
    apply_cell(ws4, row, 5, f"¥{total:,}", bold_font)
    apply_cell(ws4, row, 6, roi, normal_font)

ws4.freeze_panes = "A3"
auto_width(ws4)

# ========== 图表 Sheet ==========
ws_chart = wb.create_sheet("分数走势图")

# 柱状图
chart = BarChart()
chart.type = "col"
chart.title = "近三年录取分数对比"
chart.y_axis.title = "综合分"
chart.x_axis.title = "院校"
chart.style = 10
chart.width = 20
chart.height = 12

# 写入图表数据
ws_chart.cell(row=1, column=1, value="院校")
ws_chart.cell(row=1, column=2, value="2023")
ws_chart.cell(row=1, column=3, value="2024")
ws_chart.cell(row=1, column=4, value="2025")
ws_chart.cell(row=1, column=5, value="你的分数")

for i, t in enumerate(targets):
    ws_chart.cell(row=i+2, column=1, value=t["school"])
    ws_chart.cell(row=i+2, column=2, value=t.get("score_2023", 0))
    ws_chart.cell(row=i+2, column=3, value=t.get("score_2024", 0))
    ws_chart.cell(row=i+2, column=4, value=t.get("score_2025", 0))
    ws_chart.cell(row=i+2, column=5, value=comprehensive)

data_ref = Reference(ws_chart, min_col=2, max_col=4, min_row=1, max_row=len(targets)+1)
cats_ref = Reference(ws_chart, min_col=1, min_row=2, max_row=len(targets)+1)
chart.add_data(data_ref, titles_from_data=True)
chart.set_categories(cats_ref)

# 用户分数参考线
from openpyxl.chart.series import DataPoint
user_ref = Reference(ws_chart, min_col=5, max_col=5, min_row=1, max_row=len(targets)+1)
chart.add_data(user_ref, titles_from_data=True)

ws_chart.add_chart(chart, "A10")

# === 保存 ===
output = input_file.replace('.json', '.xlsx') if input_file else 'XF志愿填报方案.xlsx'
wb.save(output)
print(f"Done:: {output}")